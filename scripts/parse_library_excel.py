"""
parse_library_excel.py
Reads every .xlsx file in data/raw/ and combines them into data/books.csv.

Expected Excel columns (row 1, case-insensitive):
    title, author, category, subcategory, description, keywords,
    difficulty, call_number, year, toc, availability, location

Auto-added by parser: id, source
"""

import csv
import os
import re
import glob

try:
    import openpyxl
except ImportError:
    print("❌ openpyxl not installed. Run:")
    print("   pip install openpyxl")
    raise SystemExit(1)


RAW_DIR = "data/raw"
OUTPUT_CSV = "data/books.csv"
SOURCE_LABEL = "NWU ERC OPAC"

EXPECTED_COLS = [
    "title", "author", "category", "subcategory", "description",
    "keywords", "difficulty", "call_number", "year", "toc",
    "availability", "location",
]

FINAL_COLS = ["id"] + EXPECTED_COLS + ["source"]


def norm_header(h):
    """Normalize header: lowercase, spaces/hyphens -> underscores."""
    if h is None:
        return ""
    return re.sub(r"[\s\-]+", "_", str(h).strip().lower())


def norm_value(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def clean_keywords(raw):
    """Turn 'subject headings' into a clean comma-separated list."""
    text = norm_value(raw).lower()
    text = re.sub(r"\b\d+\.\s*", "", text)
    parts = re.split(r"[,;\.]+", text)
    seen, out = set(), []
    for p in parts:
        p = p.strip()
        if not p or p in seen:
            continue
        seen.add(p)
        out.append(p)
    return ",".join(out)


def clean_availability(raw):
    text = norm_value(raw).lower()
    if "borrow" in text:
        return "borrowed"
    return "available"


def clean_location(raw):
    text = norm_value(raw).lower()
    if "portal" in text or "ebook" in text or "online" in text:
        return "portal"
    return "on the shelf"


def clean_difficulty(raw):
    text = norm_value(raw).lower()
    if text in ("beginner", "intermediate", "advanced"):
        return text
    return "intermediate"


def parse_file(path):
    print(f"  📖 Reading: {os.path.basename(path)}")
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet = wb[wb.sheetnames[0]]

    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        print("     ⚠️  Empty file, skipping")
        return []

    # Find header row
    header_row_idx = None
    headers = []
    for i, row in enumerate(rows[:5]):
        normalized = [norm_header(c) for c in row]
        if "title" in normalized:
            header_row_idx = i
            headers = normalized
            break

    if header_row_idx is None:
        print("     ⚠️  No header row with 'title' found, skipping")
        return []

    col_idx = {}
    for col in EXPECTED_COLS:
        col_idx[col] = headers.index(col) if col in headers else None

    missing = [c for c, i in col_idx.items() if i is None]
    if missing:
        print(f"     ⚠️  Missing columns: {missing}")

    books = []
    for row in rows[header_row_idx + 1:]:
        if not row:
            continue
        title = norm_value(row[col_idx["title"]]) if col_idx["title"] is not None else ""
        if not title or len(title) < 3:
            continue

        book = {}
        for col in EXPECTED_COLS:
            idx = col_idx[col]
            raw = row[idx] if idx is not None and idx < len(row) else ""
            if col == "keywords":
                book[col] = clean_keywords(raw)
            elif col == "availability":
                book[col] = clean_availability(raw)
            elif col == "location":
                book[col] = clean_location(raw)
            elif col == "difficulty":
                book[col] = clean_difficulty(raw)
            else:
                book[col] = norm_value(raw)
        books.append(book)

    print(f"     ✅ {len(books)} books loaded")
    return books


def main():
    if not os.path.isdir(RAW_DIR):
        print(f"❌ Folder {RAW_DIR} does not exist.")
        print(f"   Create it and put your .xlsx files inside.")
        return

    files = sorted(glob.glob(os.path.join(RAW_DIR, "*.xlsx")))
    files = [f for f in files if not os.path.basename(f).startswith("~$")]
    if not files:
        print(f"❌ No .xlsx files found in {RAW_DIR}/")
        return

    print(f"📚 Found {len(files)} Excel file(s) in {RAW_DIR}/\n")

    all_books = []
    for f in files:
        all_books.extend(parse_file(f))

    if not all_books:
        print("\n❌ No books parsed. Check the Excel header row.")
        return

    for i, b in enumerate(all_books, start=1):
        b["id"] = i
        b["source"] = SOURCE_LABEL

    os.makedirs(os.path.dirname(OUTPUT_CSV) or ".", exist_ok=True)
    with open(OUTPUT_CSV, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=FINAL_COLS)
        w.writeheader()
        w.writerows(all_books)

    print(f"\n✅ Wrote {len(all_books)} books to {OUTPUT_CSV}")
    print(f"📋 Columns: {', '.join(FINAL_COLS)}")


if __name__ == "__main__":
    main()